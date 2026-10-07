"""Bounded, streaming readers. Import rows remain lossless strings until quality processing."""

from __future__ import annotations

import csv
import zipfile
from datetime import date, datetime
from pathlib import Path

from .settings import MAX_CELL_CHARS, MAX_COLUMNS, MAX_XLSX_BYTES


class DataIssue(Exception):
    def __init__(self, message: str, code: str = "invalid_file", status: int = 422):
        super().__init__(message)
        self.message, self.code, self.status = message, code, status


def inspect_xlsx(path: Path) -> dict:
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            if len(infos) > 10000 or sum(i.file_size for i in infos) > MAX_XLSX_BYTES:
                raise DataIssue("工作簿解压后超过 1 GiB，请缩小文件。", "xlsx_limit", 413)
            if any(i.flag_bits & 1 for i in infos):
                raise DataIssue("请先解除工作簿加密，再上传。")
            if any("vbaproject" in i.filename.lower() for i in infos):
                raise DataIssue("请另存为不含宏的 XLSX 文件。")
            if "xl/workbook.xml" not in z.namelist():
                raise DataIssue("文件不是有效的 XLSX 工作簿。")
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            return {"kind": "xlsx", "sheets": book.sheetnames}
        finally:
            book.close()
    except DataIssue:
        raise
    except Exception as exc:
        raise DataIssue("无法读取工作簿，请确认文件完整且没有加密。") from exc


def text_value(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    value = str(value)
    if len(value) > MAX_CELL_CHARS:
        raise DataIssue("存在超过 32,768 字符的单元格，请缩小内容。", "cell_limit")
    return value if value.strip() else None


def headers(values) -> list[str]:
    if len(values) == 0 or len(values) > MAX_COLUMNS:
        raise DataIssue("表格需包含 1–200 列。", "column_limit", 413)
    names = [str(v).strip() if v is not None else "" for v in values]
    if any(not n or len(n) > 200 for n in names):
        raise DataIssue("请为每列提供非空且不超过 200 字符的表头。", "invalid_headers")
    if len(set(names)) != len(names):
        raise DataIssue("表头重复，请先为各列提供不同名称。", "duplicate_headers")
    return names


def iter_csv(path: Path, options: dict):
    csv.field_size_limit(MAX_CELL_CHARS)
    try:
        with path.open(encoding=options["encoding"], newline="", errors="strict") as f:
            reader = csv.reader(f, delimiter=options["delimiter"], strict=True)
            try:
                names = headers(next(reader))
            except StopIteration as exc:
                raise DataIssue("文件为空，请选择包含表头和数据的文件。") from exc
            yield "headers", names
            for record, row in enumerate(reader, start=2):
                if not row or all(not v.strip() for v in row):
                    yield "blank", None
                    continue
                if len(row) != len(names):
                    raise DataIssue(f"第 {record} 条记录的列数与表头不一致，请检查分隔符。")
                yield "row", ("CSV", record, [text_value(v) for v in row])
    except UnicodeError as exc:
        raise DataIssue(
            "当前编码无法读取文件，请选择正确编码后重新预览。", "encoding_error"
        ) from exc
    except csv.Error as exc:
        raise DataIssue("CSV 引号、分隔符或字段长度不符合要求，请检查文件。") from exc


def iter_xlsx(path: Path, options: dict):
    inspection = inspect_xlsx(path)
    chosen = options["sheets"]
    if not chosen or len(chosen) != len(set(chosen)) or set(chosen) - set(inspection["sheets"]):
        raise DataIssue("请选择一个或多个有效且不重复的工作表。", "invalid_sheets")
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    expected = None
    try:
        for sheet in chosen:
            ws = book[sheet]
            if ws.max_column and ws.max_column > MAX_COLUMNS:
                raise DataIssue("工作表超过 200 列，请缩小范围。", "column_limit", 413)
            rows = ws.iter_rows(values_only=True)
            first = next(rows, None)
            if first is None:
                raise DataIssue(f"工作表「{sheet}」为空。")
            names = headers(first)
            if expected is None:
                expected = names
                yield "headers", names
            elif names != expected:
                raise DataIssue(
                    "合并的工作表必须具有相同名称和顺序的表头。", "sheet_schema_mismatch"
                )
            for number, row in enumerate(rows, start=2):
                if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
                    yield "blank", None
                    continue
                if len(row) > len(names):
                    raise DataIssue(f"工作表「{sheet}」第 {number} 行包含无表头的数据。")
                values = list(row) + [None] * (len(names) - len(row))
                yield "row", (sheet, number, [text_value(v) for v in values])
    finally:
        book.close()
